"""Builds the Group 31 project report (Word) from the repository's numbers."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
fea = json.load(open(os.path.join(ROOT, "hardware", "fea", "fea_results.json")))

doc = Document()
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)

def font(style, size=12, bold=None):
    f = style.font
    f.name = "Times New Roman"; f.size = Pt(size); f.color.rgb = RGBColor(0, 0, 0); f.italic = False
    if bold is not None: f.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")

font(doc.styles["Normal"], 12)
for name, size in (("Title", 12), ("Heading 1", 12), ("Heading 2", 12), ("Heading 3", 12), ("Caption", 12)):
    font(doc.styles[name], size, bold=(name != "Caption"))
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15

fig_n = [0]; tab_n = [0]

def H(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Times New Roman"; r.font.size = Pt(12); r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0); r.font.italic = False
    return p

def P(text, bold=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = False
    if align: p.alignment = align
    return p

def bullets(items):
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(t).italic = False

def caption(text, kind):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text).italic = False
    return p

def figure(path, text, width=15):
    fig_n[0] += 1
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(f"Figure {fig_n[0]}: {text}", "fig")
    return fig_n[0]

def table(header, rows, text, widths=None):
    tab_n[0] += 1
    caption(f"Table {tab_n[0]}: {text}", "tab")
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return tab_n[0]

# ---------------------------------------------------------------- title page
P("MECN3066A Mechatronics II", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Project Report: Recycled Robotic Arm", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Group 31", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Group members and student numbers: [to be completed by the group]", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Demonstration: Monday 26 October 2026, NWE Laboratory", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Submitted: [date]", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------------------------------------------------------- 1 summary
H("1. Summary")
P("This report describes a three degree of freedom robotic arm built mainly from recycled foam board and plywood, driven by four SG90 micro servos and controlled by an Arduino Uno. The arm works inside the 400 mm half circle set by the brief and completes four tasks without intervention: it checks Points A and B for an object with an ultrasonic sensor and answers with one of three gestures, reads the colour at Point C and lights a matching LED, measures the loudness of a tap at Point D and points at a matching height on a volume indicator, and presses a three letter word in Morse code on a push button at Point E, which a second Uno decodes and displays.")
P("Joint angles are calculated on the board from the station coordinates with two link inverse kinematics that account for where each sensor sits on the sensor head. A finite element analysis of both links shows peak stresses below 0.1 MPa under twice the static load, so the links are not the limiting part. A torque check shows the opposite for the actuators: holding the arm at Point B needs 86% of the stall torque of the two shoulder servos, and the worst reachable pose needs 95%. The tension spring on the main arm and the gear mesh at the base are therefore the parts of the design that decide whether the arm works reliably.")

# ---------------------------------------------------------------- 2 brief
H("2. Brief and requirements")
P("The brief defines a working area of a 400 mm half circle, 400 mm high, with stations given in cylindrical coordinates (angle from the arm's left, radius from the base axis, height above the table). Table 1 lists them.")
table(["Station", "Angle (deg)", "Radius (mm)", "Height (mm)", "Task"],
      [["A", 10, 350, 320, "Object present? (ultrasonic)"],
       ["B", 10, 350, 150, "Object present? (ultrasonic)"],
       ["C", 55, 200, 70, "Colour of the target, shown on an LED"],
       ["D", 100, 275, 250, "Loudness of a tap"],
       ["Volume indicator", 110, 275, 250, "Point at a height that matches the loudness"],
       ["E", 170, 250, 300, "Press a word in Morse code on a push button"]],
      "Stations from the project brief")
P("Other requirements from the brief and the demonstration notice:")
bullets(["The button at Point E sits on a breadboard that the tutor mounts on the stand; its orientation is the group's choice.",
         "The Morse receiver runs on a second Uno, separate from the arm controller, and shows the word on an LCD.",
         "Each attempt at the demonstration is filmed, and the arm controller code, the receiver code and the videos are uploaded straight after the slot.",
         "The arm runs from reset with no help during an attempt. The best of the attempts in the 10 minute slot counts."])

# ---------------------------------------------------------------- 3 mechanical
H("3. Mechanical design")
H("3.1 Layout", 2)
P("The arm has a turntable base and two links in one vertical plane. The base servo turns a hub on a 20 mm plywood platform through a gear on the hub's rim. Two side plates on the hub carry the shoulder pivot 130 mm above the table, where two SG90 servos face each other and drive the main arm together in opposite directions, which doubles the shoulder torque. The second arm servo sits at the top of the main arm. A tension spring runs from the base to the main arm and carries part of the arm's weight. The sensor head at the end of the second arm holds the ultrasonic sensor on the arm axis, the sound sensor 42 mm further out and 45 mm to one side, the colour sensor 83 mm out and 45 mm to the other side, and a 70 mm button pusher hanging below the arm end.")
figure(os.path.join(HERE, "cad_rest_pose.png"), "CAD model at the rest position: base facing forward, main arm upright, second arm level (all servos at 90 degrees)", 12)
table(["Part", "Material", "Size (mm)", "Mass (g)"],
      [["Base platform", "Plywood", "250 x 250 x 20", "about 700"],
       ["Turntable hub", "Recycled tube, foam board", "120 diameter, 60 high", "about 60"],
       ["Main arm", "5 mm foam board box", "225 pivot to pivot, 50 x 30 section", f"{fea['main_arm']['link_mass_g']:.0f}"],
       ["Second arm", "4 mm foam board box with three windows", "160 to the sensor head, 40 x 20 section", f"{fea['second_arm_case_A']['link_mass_g']:.0f}"],
       ["Sensor head", "Plywood, sensors, pusher, wiring", "110 wide", "46 (estimate)"],
       ["Servos", "SG90, 9 g, 1.8 kgf cm stall at 4.8 V", "23 x 12.2 x 22.7", "4 x 9"]],
      "Main parts. Link masses come from the CAD volume and an assumed board density of 500 kg/m3")
H("3.2 Materials", 2)
P("The links are boxes made from white 5 mm foam board taped at the edges, with windows cut in the second arm to save weight. The platform and the sensor head are offcut plywood. The material properties used in the analysis are those of expanded PVC foam board (elastic modulus 800 MPa, Poisson's ratio 0.35, density 500 kg/m3, tensile strength 15 MPa), because the exact grade of the recycled board is not known. Section 6 shows that the stresses are so low that this uncertainty does not change any conclusion.")
H("3.3 CAD files", 2)
P("The model is parametric (CadQuery, file hardware/cad/arm/build_cad.py) and uses the same dimensions as the controller code. It exports a STEP file for every part and two assemblies, one at the rest position and one at Point B. Figure 2 shows two of the working poses.")
figure(os.path.join(HERE, "cad_point_B.png"), "CAD model at Point B: base 80 degrees left, main arm 46.6 degrees forward, second arm 11 degrees up", 11)

# ---------------------------------------------------------------- 4 kinematics
H("4. Kinematics")
P("All joint angles are relative to the rest position. The base angle is 90 degrees minus the station angle, corrected for the sideways offset of the sensor in use. In the arm's vertical plane the main arm is link 1 (225 mm, angle phi from vertical) and the second arm plus the sensor's offset along it is link 2 (160 mm plus the offset). For a target at in-plane distance r and height z, the sensor must sit a set standoff s back from the target along the direction it points. That point is solved with the law of cosines (elbow up), and because the pointing direction depends on the answer, the solution is repeated six times, which converges well within 0.1 mm.")
P("For the button press the target is the pusher tip itself, with no standoff, so the arm can drive the tip along any line. Each press moves the tip straight down from 36 mm above the button to 12 mm into it. Pushing past contact means the servos keep pressing hard for the whole hold time.")
table(["Point", "Sensor", "Standoff (mm)", "Base (deg)", "Main arm (deg)", "Second arm (deg)"],
      [["A", "Ultrasonic", 100, "80.0 left", "23.7 forward", "20.4 up"],
       ["B", "Ultrasonic", 70, "80.0 left", "46.6 forward", "11.0 up"],
       ["C", "Colour", 15, "22.0 left", "33.3 forward", "40.6 down"],
       ["D", "Sound", 40, "0.6 right", "13.0 forward", "11.2 down"],
       ["Indicator, 150 mm", "Ultrasonic", 60, "20.0 right", "34.9 forward", "13.3 down"],
       ["Indicator, 350 mm", "Ultrasonic", 60, "20.0 right", "14.2 forward", "14.6 up"],
       ["E, above the button", "Pusher tip", "36 above", "80.0 right", "17.2 forward", "43.8 up"],
       ["E, pressing", "Pusher tip", "12 in", "80.0 right", "19.4 forward", "31.5 up"]],
      "Joint angles calculated by the controller for each point (relative to the rest position)")
P("Point C on the stands is a vertical blue face 70 mm above the table. With these link lengths and a shoulder 130 mm high the arm cannot bring the colour sensor down level with it: at C the second arm always points about 75 degrees down, so the sensor looks steeply down at the face from a few millimetres away. The code supports a colour sensor tilted up on a wedge (COLOUR_PITCH_DEG); with a 60 to 70 degree wedge the sensor looks at the face almost level while the arm pose hardly changes.")

# ---------------------------------------------------------------- 5 actuators
H("5. Actuator torque")
P("The holding torque at each joint was calculated from the link masses in Table 2 with the second arm at its working angle. Table 4 compares it with the stall torque of the SG90 (1.8 kgf cm at 4.8 V, two servos at the shoulder).")
table(["Pose", "Shoulder (kgf cm)", "Share of 2 x SG90", "Elbow (kgf cm)", "Share of 1 x SG90"],
      [["A", "2.28", "63%", "1.03", "57%"],
       ["B", "3.10", "86%", "0.84", "46%"],
       ["C", "2.00", "55%", "0.29", "16%"],
       ["D", "1.64", "45%", "0.94", "52%"],
       ["E, above the button", "1.84", "51%", "0.92", "51%"],
       ["Worst reachable: main arm 50 deg, second arm level", "3.41", "95%", "1.03", "57%"]],
      "Static holding torque without the spring, against servo stall torque")
P("Micro servos lose accuracy and heat up well before stall, so the shoulder relies on the tension spring. Without it, Point B and any pose with the main arm leaning far forward would sag or stall. The elbow servo runs at half its stall torque just holding the head, which leaves little in reserve for pressing: a typical 6 mm tactile switch needs 1.6 to 2.5 N, and the elbow alone can add about 1 N at the 184 mm pusher lever. The straight-line press therefore uses the shoulder and elbow together and drives 12 mm past contact so both servos push at close to full torque during the hold.")
P("The base servo carries the turning load through a gear on the rim of the hub. When the arm leans out, the reaction lifts the hub on one side and the gear can ride out of the groove, after which every point is off by the angle lost. The controller now starts each servo gently (Section 7.3), but the lasting fix is mechanical: a washer or lip that keeps the gear in the groove, a stiffer servo mount, and if possible a thrust bearing under the hub so the base servo carries none of the arm's weight.")

# ---------------------------------------------------------------- 6 FEA
H("6. Structural analysis (FEA)")
H("6.1 Method", 2)
P("Each link was meshed with 8-node hexahedral elements on a regular 2 to 2.5 mm grid that follows its real section: a hollow box with 5 mm walls for the main arm, and a 4 mm walled box with the three side windows for the second arm. Each link is clamped over the region gripped by the servo horn. Loads are the link's own weight, the weight of everything it carries and, for the second arm, the button-press reaction. A dynamic factor of 2 is applied to all weights to cover the accelerations of fast moves and the start-up jerk. The solver is written in Python (hardware/fea/fea_links.py) and solves the sparse stiffness system directly.")
v = fea["verification"]
P(f"Before use, the solver was checked on a solid 20 x 20 x 200 mm cantilever with a 1 N tip load. The tip deflection was {abs(v['tip_fe']):.4f} mm against {abs(v['tip_beam']):.4f} mm from beam theory with shear ({v['err_pct']:+.1f}%). A pure end moment on the main arm section and a tip load on the same section each gave bending stresses within 6% of M c / I at mid span, which also confirmed the sign of the elbow moment.")
H("6.2 Load cases and results", 2)
a, b, m = fea["second_arm_case_A"], fea["second_arm_case_B"], fea["main_arm"]
ld = fea["loads"]
table(["Case", "Elements", "Peak von Mises (MPa)", "Safety factor on 15 MPa", "Tip deflection (mm)"],
      [["Second arm level, weights x2 (head 46 g at the tip)", a["elements"], a["max_von_mises_MPa"], a["safety_factor"], a["tip_deflection_mm"]],
       ["Second arm pressing, weight x1 with a 0.96 N press reaction", b["elements"], b["max_von_mises_MPa"], b["safety_factor"], b["tip_deflection_mm"]],
       [f"Main arm at 50 deg, elbow force {ld['elbow_force_N']} N and moment {ld['elbow_moment_Nmm']} N mm, weights x2", m["elements"], m["max_von_mises_MPa"], m["safety_factor"], m["tip_deflection_mm"]]],
      "FEA results (peak stress outside the clamp and load patches)")
figure(os.path.join(ROOT, "hardware", "fea", "fea_second_arm.png"), "Second arm: von Mises stress with all weights doubled. The highest stresses sit at the corners of the windows", 16)
figure(os.path.join(ROOT, "hardware", "fea", "fea_main_arm.png"), "Main arm at its 50 degree limit: von Mises stress, highest in the outer walls next to the shoulder horn", 16)
H("6.3 Discussion", 2)
P("Peak stresses are below 0.1 MPa, at least 150 times below the strength of the board, and tip deflections are below 0.1 mm. Hand calculation agrees: the main arm moment at the shoulder in the worst pose is about 670 N mm, giving 0.08 MPa in the outer wall against 0.075 MPa from the FEA. The window corners in the second arm raise the local stress by about 50% over the plain section, which is still negligible. The links could be made lighter, and the useful saving is mass rather than strength, since every gram at the head costs elbow and shoulder torque (Section 5). The FEA does not cover the taped joints, the servo horn screws or the gear mesh, which are the parts most likely to fail in practice.")

# ---------------------------------------------------------------- 7 electronics & software
H("7. Electronics and software")
H("7.1 Electronics", 2)
table(["Device", "Pin", "Notes"],
      [["Base servo", "D2", "SG90"], ["Main arm servos, left and right", "D3, D4", "Driven in opposite directions"],
       ["Second arm servo", "D5", "SG90"], ["Ultrasonic HC-SR04", "Echo D6, Trig D7", "5 V"],
       ["Sound sensor", "A1", "Analogue output"], ["Colour sensor TCS34725", "SDA A4, SCL A5", "Hardware I2C"],
       ["Red, green, blue LEDs", "D8, D9, D10", "220 ohm each, active high"], ["Beeper", "D11", "Passive, 2.5 kHz tone"],
       ["LCD monitor 16 x 2, I2C backpack", "SDA A2, SCL A3", "Software I2C, address 0x27 or 0x3F"],
       ["Servo supply", "External 5 V, 2 A", "Common ground with the Uno"],
       ["Morse receiver (second Uno)", "Button D7, LCD on A4/A5", "Separate board"]],
      "Arm controller pin map")
P("The servos run from an external 5 V supply because four SG90s can draw more than an amp together when they start or stall, far more than the Uno's regulator can give. The LCD monitor uses A2 and A3 because the colour sensor already occupies the hardware I2C pins; a small software I2C driver in the sketch talks to it, so no library is needed.")
H("7.2 Program structure", 2)
P("The arm controller is a state machine: rest hold, Task 1, Task 2, Task 3, Task 4, finish, idle. During the 5 s rest hold the Morse word can be typed in the Serial Monitor; it is saved in EEPROM. Each task beeps its number at the start and at the end and then waits 2 s. Every move uses cosine easing with a top joint speed of 30 degrees per second, which keeps the motion smooth and avoids jerks. The Serial Monitor prints every move in plain words, and the LCD shows the task, the distances at A and B, each result and a summary at the end.")
bullets(["Task 1: the arm goes to A and measures the distance (median of five pings, at least three echoes). If nothing is straight ahead it looks a little left, right, up and down. An object counts between 2 and 25 cm. It then lifts and turns away so the visit to B is clearly separate. Object at A: YES (the sensor moves straight up and down). Object at B: NO (the base turns left and right). Nothing at either: BOX (the sensor traces a square).",
         "Task 2: at C the colour sensor waves over five spots, the LEDs follow what it sees, and the clearest reading decides the colour. The matching LED stays lit to the end of the run.",
         "Task 3: at D the arm waits up to 3 s for a tap and measures its peak to peak amplitude. It then moves to the bottom of the indicator scale and rises slowly to a height from 150 mm (quiet) to 350 mm (loud), drops back and rises again like a level meter, and holds.",
         "Task 4: at E each press drives the pusher straight down and up at 120 degrees per second. Dots are 600 ms and dashes 1800 ms, with 600 ms between presses and 1800 ms between letters."])
H("7.3 Start-up and calibration", 2)
P("A hobby servo cannot report its position, so the first pulse after power-up sends it to the commanded angle at full speed. On this arm that jerk pulled cables loose and lifted the base gear out of its groove. The controller now sends each servo short bursts of pulses with pauses between them, so it creeps to the rest position: second arm first, then the main arm pair together, then the base, slowest of all. Two calibration sketches support the set-up. test_servos lets each joint be nudged to the true rest position (the two main arm servos can be trimmed separately so they do not fight) and prints the zero values to paste into the code. capture_points takes the arm to every station as the demonstration does, lets each one be nudged onto the real stand, tunes the press depth and height at E, and prints the corrected station lines.")
H("7.4 Morse timing", 2)
P("The receiver splits dots from dashes at 1250 ms and ends a letter after 1200 ms of silence. A release shorter than 100 ms is treated as contact bounce within the same press. The press timing was checked in an offline model of the arm that runs the real sketch, tracks the pusher tip from the servo commands and records when the button would close. With the button from 30 mm higher to 12 mm lower than its nominal position, the button saw dots of 520 to 800 ms, dashes of 1720 to 2000 ms, gaps of 400 to 800 ms inside a letter and 1600 to 1880 ms between letters, so every band is at least 400 ms from a threshold. Both receiver sketches were then run against these recorded contact timelines: 1,536 runs covering eight words, eight button heights and contact chatter up to 90 ms, and 600 runs with every edge shifted randomly by up to 150 ms, all decoded correctly.")

# ---------------------------------------------------------------- 8 testing
H("8. Testing and results")
H("8.1 Tests before the demonstration", 2)
bullets(["Every sketch is compiled for the Uno on each change by an automated build (GitHub Actions).",
         "A browser simulator runs the same joint maths and task sequence and draws the arm to scale (simulation/arm_bench.html). It was used to check reach, joint limits and the order of moves.",
         "An offline model compiles the real sketch on a PC with simulated sensors, servos, LCD and button, and checks the beep pattern, gestures, joint limits, LCD text and Morse decoding for different conditions.",
         "On the real arm, test_servos, test_sensors and test_combined were used to set the zero positions, directions, colour references and sound levels, and to visit every station."])
H("8.2 Problems found and fixes", 2)
table(["Problem on the real arm", "Cause", "Fix"],
      [["Main arm leaned back when it should lean forward", "Servo direction", "MAIN_DIR set to -1"],
       ["Points off after start-up", "Zero positions not set, base gear jumping teeth", "Zero calibration in test_servos; soft start; gear retention"],
       ["Base gear lifts out of the turntable groove", "Start-up jerk and the arm's moment on the hub", "Soft start, base last and slowest; mechanical retention recommended"],
       ["Button presses too weak", "Elbow servo near half its stall torque just holding the head", "Straight-line press with both joints, 12 mm past contact, 120 deg/s"],
       ["Visit to B not visible", "A and B share one bearing", "Lift and turn away between A and B"],
       ["Point C is a vertical face", "Arm cannot reach low enough to look level", "Sensor waves over five spots; optional tilted sensor mount"],
       ["Button on a sloping stand", "Stand design", "Breadboard mounted so the button faces up; straight-down press"]],
      "Problems found during testing")
H("8.3 Demonstration results", 2)
P("[To be completed after the demonstration: for each attempt, the conditions set by the tutors (object at A, B or neither; colour; tap; word), what the arm did at each point, and whether each task passed. Link the videos.]")
table(["Attempt", "Task 1 (A/B, gesture)", "Task 2 (colour)", "Task 3 (sound)", "Task 4 (Morse)"],
      [["1", "", "", "", ""], ["2", "", "", "", ""], ["3", "", "", "", ""]], "Demonstration results (to be filled in)")

# ---------------------------------------------------------------- 9 recommendations
H("9. Recommendations")
bullets(["Hold the base gear in mesh with a lip or washer and carry the arm's weight on a thrust bearing, so the base servo only turns the hub.",
         "Replace the shoulder SG90 pair with metal gear servos such as the MG90S (2.2 kgf cm) or a single MG996R, or keep the spring and check its preload so the shoulder runs well below stall at Point B.",
         "Lighten the sensor head: every 10 g there costs about 0.16 kgf cm at the elbow.",
         "Mount the colour sensor on a 60 to 70 degree wedge so it faces the vertical target at C squarely.",
         "Use a low-force button (below 1.5 N) at Point E, or add a short spring-loaded tip to the pusher so presses are consistent."])

# ---------------------------------------------------------------- 10 conclusion
H("10. Conclusion")
P("The arm reaches every station in the brief with the right sensor at a set distance, using joint angles calculated on the board. The foam board links are far stronger than they need to be, while the servos, the spring and the base gear set the limits of the design. The changes made during testing (soft start, calibrated zero positions, a firm straight-line press, a wider search at A and B, a colour scan at C and a clear level-meter motion at D) target the failures seen on the real arm.")

# ---------------------------------------------------------------- references
H("References")
for r in ["MECN3066A Mechatronics II Project Brief 2026, University of the Witwatersrand.",
          "TowerPro SG90 micro servo datasheet.",
          "ams AG, TCS3472 colour light-to-digital converter datasheet.",
          "Elecfreaks, HC-SR04 ultrasonic ranging module user guide.",
          "Hitachi, HD44780U dot matrix LCD controller datasheet; NXP, PCF8574 remote 8-bit I/O expander datasheet.",
          "Arduino, Servo library reference, https://www.arduino.cc/reference/en/libraries/servo/.",
          "R. C. Hibbeler, Mechanics of Materials, Pearson."]:
    p = doc.add_paragraph(style="List Number"); p.add_run(r)

# ---------------------------------------------------------------- appendix
H("Appendix A: Files")
table(["File", "Contents"],
      [["arduino/arm_controller/arm_controller.ino", "Demonstration program"],
       ["arduino/morse_receiver, morse_receiver_i2c", "Receiver on the second Uno"],
       ["arduino/tests/test_servos", "Zero positions, directions, sweep, manual control"],
       ["arduino/tests/test_sensors", "Sensor readings and calibration"],
       ["arduino/tests/test_combined", "Every station and movement without decisions"],
       ["arduino/tests/capture_points", "Station trims and press tuning on the real stands"],
       ["hardware/cad/arm/*.step", "CAD parts and assemblies"],
       ["hardware/fea/fea_links.py, fea_results.json", "FEA solver, load cases and results"],
       ["simulation/arm_bench.html", "Browser simulator"],
       ["docs/demo_day.md, docs/wiring.md", "Demonstration procedure and wiring"]], "Files submitted with this report")

out = os.path.join(HERE, "Group31_Robotic_Arm_Report.docx")
doc.save(out)
print("saved", out)
